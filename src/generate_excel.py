"""
generate_excel.py
=================
Génère le classeur Excel "Justificatifs de Vêture – Allocation Mensuelle".

Feuilles créées :
  - Configuration : informations fixes (AF, ville) + tableau enfants/référents
  - Formulaire    : formulaire mensuel avec liste déroulante enfant et date auto

Usage :
    python src/generate_excel.py [--output output/justificatif_veture.xlsx]
"""

from __future__ import annotations

import argparse
import os
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import (
    Alignment,
    Border,
    Font,
    PatternFill,
    Side,
)
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

# ---------------------------------------------------------------------------
# Couleurs / styles
# ---------------------------------------------------------------------------
BLEU_EN_TETE = "1F4E79"   # bleu foncé pour les en-têtes
BLEU_CLAIR   = "D6E4F0"   # fond clair pour les sections
GRIS_CLAIR   = "F2F2F2"   # fond alterné
BLANC        = "FFFFFF"
JAUNE_SAISIE = "FFFACD"   # fond des cellules à remplir

def thin_border(top=True, bottom=True, left=True, right=True) -> Border:
    thin = Side(style="thin")
    return Border(
        top=thin if top else None,
        bottom=thin if bottom else None,
        left=thin if left else None,
        right=thin if right else None,
    )

def header_fill(color: str = BLEU_EN_TETE) -> PatternFill:
    return PatternFill("solid", fgColor=color)

def cell_fill(color: str) -> PatternFill:
    return PatternFill("solid", fgColor=color)


# ---------------------------------------------------------------------------
# Onglet Configuration
# ---------------------------------------------------------------------------

# Données d'exemple (modifiez ici ou directement dans l'onglet Configuration)
CONFIG_AF = {
    "nom_prenom": "Dupont Marie",
    "adresse": "12 Rue des Lilas, 09000 Foix",
    "telephone": "06 12 34 56 78",
    "email": "marie.dupont@email.fr",
    "fait_a": "Foix",
}

# Exemple de mapping enfants → référents sociaux
ENFANTS_EXEMPLE = [
    ("Martin Lucas",  "Mme Bernard Sophie"),
    ("Petit Emma",    "M. Leclerc Paul"),
]


def build_configuration_sheet(wb: Workbook) -> None:
    """Crée et remplit l'onglet 'Configuration'."""
    ws = wb.create_sheet("Configuration")

    # ---- Largeurs de colonnes ------------------------------------------
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 40
    ws.column_dimensions["C"].width = 5

    # ---- Titre ---------------------------------------------------------
    ws.merge_cells("A1:B1")
    c = ws["A1"]
    c.value = "⚙  CONFIGURATION"
    c.font = Font(bold=True, size=14, color=BLANC)
    c.fill = header_fill(BLEU_EN_TETE)
    c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28

    ws.row_dimensions[2].height = 6  # spacer

    # ---- Sous-titre informations fixes ---------------------------------
    ws.merge_cells("A3:B3")
    c = ws["A3"]
    c.value = "Informations fixes – Assistant(e) Familial(e)"
    c.font = Font(bold=True, size=11, color=BLANC)
    c.fill = header_fill("2E75B6")
    c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[3].height = 22

    fields = [
        ("Nom / Prénom AF",        CONFIG_AF["nom_prenom"]),
        ("Adresse",                CONFIG_AF["adresse"]),
        ("Téléphone",              CONFIG_AF["telephone"]),
        ("Email",                  CONFIG_AF["email"]),
        ("Fait à (ville)",         CONFIG_AF["fait_a"]),
    ]
    for i, (label, value) in enumerate(fields):
        row = 4 + i
        lbl = ws.cell(row=row, column=1, value=label)
        lbl.font = Font(bold=True, size=10)
        lbl.fill = cell_fill(BLEU_CLAIR)
        lbl.border = thin_border()
        lbl.alignment = Alignment(vertical="center", indent=1)

        val = ws.cell(row=row, column=2, value=value)
        val.font = Font(size=10)
        val.fill = cell_fill(JAUNE_SAISIE)
        val.border = thin_border()
        val.alignment = Alignment(vertical="center", indent=1)
        ws.row_dimensions[row].height = 18

    ws.row_dimensions[9].height = 8  # spacer

    # ---- Sous-titre tableau enfants ------------------------------------
    ws.merge_cells("A10:B10")
    c = ws["A10"]
    c.value = "Enfants accueillis et référents sociaux"
    c.font = Font(bold=True, size=11, color=BLANC)
    c.fill = header_fill("2E75B6")
    c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[10].height = 22

    # En-têtes du tableau
    headers = ["Enfant (Nom Prénom)", "Référent social"]
    for col, h in enumerate(headers, start=1):
        cell = ws.cell(row=11, column=col, value=h)
        cell.font = Font(bold=True, size=10, color=BLANC)
        cell.fill = header_fill(BLEU_EN_TETE)
        cell.border = thin_border()
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[11].height = 20

    # Données exemples
    for i, (enfant, referent) in enumerate(ENFANTS_EXEMPLE):
        row = 12 + i
        c_e = ws.cell(row=row, column=1, value=enfant)
        c_r = ws.cell(row=row, column=2, value=referent)
        fill = cell_fill(BLANC if i % 2 == 0 else GRIS_CLAIR)
        for c in (c_e, c_r):
            c.fill = fill
            c.border = thin_border()
            c.alignment = Alignment(vertical="center", indent=1)
            c.font = Font(size=10)
        ws.row_dimensions[row].height = 18

    # Lignes vides supplémentaires pour ajouts futurs
    for row in range(12 + len(ENFANTS_EXEMPLE), 30):
        for col in (1, 2):
            c = ws.cell(row=row, column=col, value="")
            c.fill = cell_fill(BLANC if (row % 2 == 0) else GRIS_CLAIR)
            c.border = thin_border()
            c.alignment = Alignment(vertical="center", indent=1)
        ws.row_dimensions[row].height = 18

    # Note bas de page
    ws.merge_cells("A31:B31")
    note = ws["A31"]
    note.value = (
        "ℹ  Ajoutez autant de lignes enfant/référent que nécessaire (lignes 12 à 29)."
    )
    note.font = Font(italic=True, size=9, color="666666")
    note.alignment = Alignment(horizontal="left", indent=1)

    # Protège les libellés (optionnel – décommentez si souhaité)
    # ws.protection.sheet = True


# ---------------------------------------------------------------------------
# Onglet Formulaire
# ---------------------------------------------------------------------------

def build_formulaire_sheet(wb: Workbook) -> None:
    """Crée et remplit l'onglet 'Formulaire' avec formules dynamiques."""
    ws = wb.create_sheet("Formulaire", 0)  # premier onglet

    # ---- Largeurs de colonnes ------------------------------------------
    # A   B      C      D      E
    # N°  Label  Valeur Date   Montant
    col_widths = {"A": 8, "B": 30, "C": 35, "D": 15, "E": 14}
    for col, w in col_widths.items():
        ws.column_dimensions[col].width = w

    # ---- Ligne 1 : Titre principal -------------------------------------
    ws.merge_cells("A1:E1")
    t = ws["A1"]
    t.value = "JUSTIFICATIFS DE VÊTURE – ALLOCATION MENSUELLE"
    t.font = Font(bold=True, size=14, color=BLANC)
    t.fill = header_fill(BLEU_EN_TETE)
    t.alignment = Alignment(horizontal="center", vertical="center")
    t.border = thin_border()
    ws.row_dimensions[1].height = 30

    # ---- Ligne 2 : Sous-titre ------------------------------------------
    ws.merge_cells("A2:E2")
    st = ws["A2"]
    st.value = (
        "Pour être pris en charge, les justificatifs de vêture doivent être "
        "présentés mensuellement et par enfant."
    )
    st.font = Font(italic=True, size=9, color="444444")
    st.fill = cell_fill(BLEU_CLAIR)
    st.alignment = Alignment(
        horizontal="center", vertical="center", wrap_text=True
    )
    st.border = thin_border()
    ws.row_dimensions[2].height = 30

    ws.row_dimensions[3].height = 8  # spacer

    # ---- Section : Identité de l'AF ------------------------------------
    ws.merge_cells("A4:E4")
    s1 = ws["A4"]
    s1.value = "IDENTITÉ DE L'ASSISTANT(E) FAMILIAL(E)"
    s1.font = Font(bold=True, size=11, color=BLANC)
    s1.fill = header_fill("2E75B6")
    s1.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    s1.border = thin_border()
    ws.row_dimensions[4].height = 22

    af_rows = [
        (5,  "Nom / Prénom :",  "=Configuration!B4"),
        (6,  "Adresse :",       "=Configuration!B5"),
        (7,  "Téléphone :",     "=Configuration!B6"),
        (8,  "Email :",         "=Configuration!B7"),
    ]
    for row, label, formula in af_rows:
        lbl = ws.cell(row=row, column=2, value=label)
        lbl.font = Font(bold=True, size=10)
        lbl.fill = cell_fill(BLEU_CLAIR)
        lbl.border = thin_border()
        lbl.alignment = Alignment(vertical="center", indent=1)

        ws.merge_cells(f"C{row}:E{row}")
        val = ws.cell(row=row, column=3, value=formula)
        val.font = Font(size=10)
        val.fill = cell_fill(BLANC)
        val.border = thin_border()
        val.alignment = Alignment(vertical="center", indent=1)
        ws.row_dimensions[row].height = 18

    ws.row_dimensions[9].height = 8  # spacer

    # ---- Section : Identité de l'enfant accueilli ----------------------
    ws.merge_cells("A10:E10")
    s2 = ws["A10"]
    s2.value = "IDENTITÉ DE L'ENFANT ACCUEILLI ET CONCERNÉ PAR L'ALLOCATION"
    s2.font = Font(bold=True, size=11, color=BLANC)
    s2.fill = header_fill("2E75B6")
    s2.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    s2.border = thin_border()
    ws.row_dimensions[10].height = 22

    # Ligne 11 : Enfant (liste déroulante en C11)
    lbl_enfant = ws.cell(row=11, column=2, value="Enfant accueilli :")
    lbl_enfant.font = Font(bold=True, size=10)
    lbl_enfant.fill = cell_fill(BLEU_CLAIR)
    lbl_enfant.border = thin_border()
    lbl_enfant.alignment = Alignment(vertical="center", indent=1)

    ws.merge_cells("C11:E11")
    val_enfant = ws["C11"]
    val_enfant.font = Font(size=10, color="000080")
    val_enfant.fill = cell_fill(JAUNE_SAISIE)
    val_enfant.border = thin_border()
    val_enfant.alignment = Alignment(vertical="center", indent=1)
    ws.row_dimensions[11].height = 20

    # Liste déroulante : valeurs depuis Configuration!A12:A29
    dv = DataValidation(
        type="list",
        formula1="Configuration!$A$12:$A$29",
        allow_blank=True,
        showDropDown=False,   # False = affiche la flèche déroulante dans Excel
        showErrorMessage=True,
        errorTitle="Valeur invalide",
        error="Veuillez sélectionner un enfant depuis la liste.",
    )
    ws.add_data_validation(dv)
    dv.add(ws["C11"])

    # Ligne 12 : Référent social (auto via INDEX/MATCH)
    lbl_ref = ws.cell(row=12, column=2, value="Référent social :")
    lbl_ref.font = Font(bold=True, size=10)
    lbl_ref.fill = cell_fill(BLEU_CLAIR)
    lbl_ref.border = thin_border()
    lbl_ref.alignment = Alignment(vertical="center", indent=1)

    ws.merge_cells("C12:E12")
    # XLOOKUP n'est pas toujours disponible selon la version Excel → INDEX/MATCH
    val_ref = ws["C12"]
    val_ref.value = (
        '=IFERROR(INDEX(Configuration!$B$12:$B$29,'
        'MATCH(C11,Configuration!$A$12:$A$29,0)),"")'
    )
    val_ref.font = Font(size=10, italic=True, color="444444")
    val_ref.fill = cell_fill(BLANC)
    val_ref.border = thin_border()
    val_ref.alignment = Alignment(vertical="center", indent=1)
    ws.row_dimensions[12].height = 20

    ws.row_dimensions[13].height = 8  # spacer

    # ---- Tableau des factures ------------------------------------------
    ws.merge_cells("A14:E14")
    s3 = ws["A14"]
    s3.value = "DÉTAIL DES FACTURES"
    s3.font = Font(bold=True, size=11, color=BLANC)
    s3.fill = header_fill("2E75B6")
    s3.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    s3.border = thin_border()
    ws.row_dimensions[14].height = 22

    # En-têtes du tableau
    table_headers = [
        (1, "A", "N°"),
        (2, "B", "Type d'Achat"),
        (3, "C", "Fournisseurs"),
        (4, "D", "Date"),
        (5, "E", "Coût (€)"),
    ]
    for col_idx, col_letter, header in table_headers:
        cell = ws.cell(row=15, column=col_idx, value=header)
        cell.font = Font(bold=True, size=10, color=BLANC)
        cell.fill = header_fill(BLEU_EN_TETE)
        cell.border = thin_border()
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[15].height = 20

    # Lignes de saisie (10 lignes)
    for i in range(10):
        row = 16 + i
        # N°
        n = ws.cell(row=row, column=1, value=i + 1)
        n.font = Font(size=9, color="888888")
        n.fill = cell_fill(GRIS_CLAIR)
        n.border = thin_border()
        n.alignment = Alignment(horizontal="center", vertical="center")

        fill = cell_fill(BLANC if i % 2 == 0 else GRIS_CLAIR)
        for col in (2, 3, 4, 5):
            c = ws.cell(row=row, column=col, value="")
            c.fill = fill
            c.border = thin_border()
            c.alignment = Alignment(vertical="center", indent=1)
            if col == 4:  # Date → format date
                c.number_format = "DD/MM/YYYY"
            if col == 5:  # Coût → format monétaire
                c.number_format = '#,##0.00 "€"'
        ws.row_dimensions[row].height = 18

    # Ligne Total
    row_total = 26
    ws.merge_cells(f"A{row_total}:D{row_total}")
    tot_lbl = ws[f"A{row_total}"]
    tot_lbl.value = "TOTAL"
    tot_lbl.font = Font(bold=True, size=10, color=BLANC)
    tot_lbl.fill = header_fill(BLEU_EN_TETE)
    tot_lbl.border = thin_border()
    tot_lbl.alignment = Alignment(horizontal="right", vertical="center", indent=1)

    tot_val = ws.cell(row=row_total, column=5)
    tot_val.value = "=SUM(E16:E25)"
    tot_val.font = Font(bold=True, size=10)
    tot_val.fill = cell_fill(BLEU_CLAIR)
    tot_val.border = thin_border()
    tot_val.number_format = '#,##0.00 "€"'
    tot_val.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[row_total].height = 20

    ws.row_dimensions[27].height = 10  # spacer

    # ---- Zone de signature --------------------------------------------
    ws.merge_cells("A28:E28")
    s4 = ws["A28"]
    s4.value = "SIGNATURE"
    s4.font = Font(bold=True, size=11, color=BLANC)
    s4.fill = header_fill("2E75B6")
    s4.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    s4.border = thin_border()
    ws.row_dimensions[28].height = 22

    # Ligne 29 : Je soussigné(e)
    ws.merge_cells("A29:E29")
    soussigne = ws["A29"]
    soussigne.value = (
        '=CONCATENATE("Je soussigné(e) ",Configuration!B4,'
        '" certifie avoir réglé les achats ci-dessus listés.")'
    )
    soussigne.font = Font(size=10, italic=True)
    soussigne.fill = cell_fill(BLANC)
    soussigne.border = thin_border()
    soussigne.alignment = Alignment(
        horizontal="left", vertical="center", wrap_text=True, indent=1
    )
    ws.row_dimensions[29].height = 22

    ws.row_dimensions[30].height = 8  # spacer

    # Ligne 31 : Fait à | date
    lbl_fait = ws.cell(row=31, column=2, value="Fait à :")
    lbl_fait.font = Font(bold=True, size=10)
    lbl_fait.fill = cell_fill(BLEU_CLAIR)
    lbl_fait.border = thin_border()
    lbl_fait.alignment = Alignment(vertical="center", indent=1)

    val_fait = ws.cell(row=31, column=3, value="=Configuration!B8")
    val_fait.font = Font(size=10)
    val_fait.fill = cell_fill(BLANC)
    val_fait.border = thin_border()
    val_fait.alignment = Alignment(vertical="center", indent=1)

    lbl_le = ws.cell(row=31, column=4, value="Le :")
    lbl_le.font = Font(bold=True, size=10)
    lbl_le.fill = cell_fill(BLEU_CLAIR)
    lbl_le.border = thin_border()
    lbl_le.alignment = Alignment(vertical="center", indent=1)

    val_le = ws.cell(row=31, column=5, value="=TODAY()")
    val_le.font = Font(size=10)
    val_le.fill = cell_fill(BLANC)
    val_le.border = thin_border()
    val_le.number_format = "DD/MM/YYYY"
    val_le.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[31].height = 20

    # Ligne 32 : M./Mme + signature
    lbl_mmme = ws.cell(row=32, column=2, value="M. / Mme :")
    lbl_mmme.font = Font(bold=True, size=10)
    lbl_mmme.fill = cell_fill(BLEU_CLAIR)
    lbl_mmme.border = thin_border()
    lbl_mmme.alignment = Alignment(vertical="center", indent=1)

    ws.merge_cells("C32:E32")
    val_mmme = ws["C32"]
    val_mmme.value = ""
    val_mmme.fill = cell_fill(JAUNE_SAISIE)
    val_mmme.border = thin_border()
    val_mmme.alignment = Alignment(vertical="center", indent=1)
    ws.row_dimensions[32].height = 28

    ws.row_dimensions[33].height = 8

    # Ligne 34 : "Signature :"
    lbl_sig = ws.cell(row=34, column=2, value="Signature :")
    lbl_sig.font = Font(bold=True, size=10)
    lbl_sig.fill = cell_fill(BLEU_CLAIR)
    lbl_sig.border = thin_border()
    lbl_sig.alignment = Alignment(vertical="center", indent=1)

    ws.merge_cells("C34:E36")
    sig_zone = ws["C34"]
    sig_zone.fill = cell_fill(BLANC)
    sig_zone.border = thin_border()
    for row in (34, 35, 36):
        ws.row_dimensions[row].height = 20

    # ---- Figer les lignes d'en-tête (row 1-2) --------------------------
    ws.freeze_panes = "A3"

    # ---- Mise en page pour l'impression --------------------------------
    from openpyxl.worksheet.page import PageMargins
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = ws.ORIENTATION_PORTRAIT
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToHeight = 1
    ws.page_setup.fitToWidth = 1
    ws.print_area = "A1:E36"
    ws.page_margins = PageMargins(
        left=0.5, right=0.5, top=0.75, bottom=0.75,
        header=0.3, footer=0.3
    )


# ---------------------------------------------------------------------------
# Point d'entrée
# ---------------------------------------------------------------------------

def generate(output_path: str) -> None:
    os.makedirs(os.path.dirname(output_path), exist_ok=True)

    wb = Workbook()
    # Supprimer la feuille par défaut
    if "Sheet" in wb.sheetnames:
        del wb["Sheet"]

    build_formulaire_sheet(wb)
    build_configuration_sheet(wb)

    wb.save(output_path)
    print(f"✅  Classeur Excel généré : {output_path}")


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Génère le classeur Excel Justificatifs de Vêture."
    )
    parser.add_argument(
        "--output",
        default="output/justificatif_veture.xlsx",
        help="Chemin du fichier Excel généré (défaut : output/justificatif_veture.xlsx)",
    )
    args = parser.parse_args()
    generate(args.output)


if __name__ == "__main__":
    main()
